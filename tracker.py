from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent
DATA_PATH = BASE_DIR / "data" / "legislation_input.csv"
OUTPUT_DIR = BASE_DIR / "outputs"
EXCEL_OUTPUT_PATH = OUTPUT_DIR / "legislative_tracker_output.xlsx"
BRIEFING_OUTPUT_PATH = OUTPUT_DIR / "daily_legislative_briefing.txt"
ERROR_LOG_PATH = OUTPUT_DIR / "error_log.txt"

REQUIRED_COLUMNS = [
    "bill_id",
    "bill_title",
    "chamber",
    "committee",
    "status",
    "last_action_date",
    "summary_text",
    "url",
    "keywords",
]

PORT_KEYWORDS = {
    "port",
    "port authority",
    "panynj",
    "maritime",
    "shipping",
    "freight",
    "logistics",
    "cargo",
    "harbor",
    "infrastructure",
    "terminal",
    "waterfront",
    "bridge",
    "tunnel",
    "toll",
}

TITLE_SUMMARY_TERMS = [
    "transportation",
    "infrastructure funding",
    "workforce development",
    "federal authorization",
    "water resources",
    "capital plan",
    "port infrastructure development program",
    "maritime workforce",
    "supply chain",
    "offshore wind supply chain",
    "cargo facility",
    "clean ports",
    "zero-emission port",
    "governance",
    "transparency",
    "new york",
    "new jersey",
]

COMMITTEE_TERMS = [
    "Transportation and Infrastructure",
    "Appropriations",
    "Commerce",
    "Homeland Security",
    "Environment and Public Works",
    "Education and Workforce",
    "Corporations, Authorities And Commissions",
    "Transportation",
    "Labor",
    "Energy And Telecommunications",
    "Budget",
]

ACTIVE_STATUS_TERMS = {
    "floor vote",
    "passed committee",
    "passed house",
    "passed assembly",
    "passed senate",
    "reported",
    "advanced to third reading",
    "delivered to governor",
    "signed by governor",
    "enacted",
}

EXCEL_COLUMN_ORDER = [
    "priority",
    "relevance_score",
    "bill_id",
    "bill_title",
    "chamber",
    "committee",
    "status",
    "last_action_date",
    "days_since_last_action",
    "summary_text",
    "keywords",
    "url",
]

EXCEL_COLUMN_WIDTHS = {
    "priority": 18,
    "relevance_score": 16,
    "bill_id": 16,
    "bill_title": 42,
    "chamber": 22,
    "committee": 34,
    "status": 28,
    "last_action_date": 18,
    "days_since_last_action": 22,
    "summary_text": 72,
    "keywords": 42,
    "url": 58,
    "metric": 34,
    "value": 18,
    "bill_count": 14,
    "average_score": 16,
}


class DataValidationError(ValueError):
    """Raised when the legislative input CSV is missing required structure."""


def log_error(msg: str) -> None:
    """Append a timestamped error message to outputs/error_log.txt."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with ERROR_LOG_PATH.open("a", encoding="utf-8") as log_file:
        log_file.write(f"[{timestamp}] {msg}\n")


def validate_data(df: pd.DataFrame) -> None:
    """Validate that the input DataFrame has the required tracker columns."""
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in df.columns]
    if missing_columns:
        message = "Missing required columns: " + ", ".join(missing_columns)
        log_error(message)
        raise DataValidationError(message)


def load_data(path: Path | str = DATA_PATH) -> pd.DataFrame:
    """Load and validate legislative CSV data."""
    try:
        df = pd.read_csv(path)
        validate_data(df)
        df["last_action_date"] = pd.to_datetime(df["last_action_date"], errors="coerce")
        invalid_dates = df["last_action_date"].isna().sum()
        if invalid_dates:
            log_error(f"{invalid_dates} row(s) contain invalid last_action_date values.")
        return df
    except DataValidationError:
        raise
    except FileNotFoundError as exc:
        message = f"Input file not found: {path}"
        log_error(message)
        raise FileNotFoundError(message) from exc
    except Exception as exc:
        log_error(f"Failed to load data: {exc}")
        raise


def _safe_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _contains_any(text: str, terms: Iterable[str]) -> bool:
    normalized = text.lower()
    return any(term.lower() in normalized for term in terms)


def _keyword_set(raw_keywords: object) -> set[str]:
    text = _safe_text(raw_keywords)
    return {keyword.strip().lower() for keyword in text.split(",") if keyword.strip()}


def _days_since_last_action(value: object) -> int | None:
    date_value = pd.to_datetime(value, errors="coerce")
    if pd.isna(date_value):
        return None
    today = pd.Timestamp.today().normalize()
    return int((today - date_value.normalize()).days)


def compute_relevance_score(row: pd.Series) -> int:
    """Compute a 0-100 relevance score for one legislative record."""
    score = 0
    keywords = _keyword_set(row.get("keywords"))
    title = _safe_text(row.get("bill_title"))
    summary = _safe_text(row.get("summary_text"))
    committee = _safe_text(row.get("committee"))
    status = _safe_text(row.get("status"))
    keyword_text = " ".join(keywords)
    combined_policy_text = f"{title} {summary}"
    full_policy_text = f"{combined_policy_text} {keyword_text}"

    if keywords.intersection(PORT_KEYWORDS):
        score += 30

    if _contains_any(combined_policy_text, TITLE_SUMMARY_TERMS):
        score += 25

    if _contains_any(committee, COMMITTEE_TERMS):
        score += 20

    if _contains_any(full_policy_text, ["wrda", "water infrastructure"]):
        score += 15

    if _contains_any(full_policy_text, ["new york", "new jersey", "panynj", "port authority"]):
        score += 15

    if status.lower() in ACTIVE_STATUS_TERMS:
        score += 10

    days_since_action = _days_since_last_action(row.get("last_action_date"))
    if status.lower() == "stalled" or (days_since_action is not None and days_since_action > 90):
        score -= 20

    return max(0, min(100, score))


def assign_priority(score: int | float) -> str:
    """Convert a relevance score into a policy priority label."""
    if score >= 80:
        return "HIGH PRIORITY"
    if score >= 50:
        return "MEDIUM PRIORITY"
    return "LOW PRIORITY"


def process_legislation(df: pd.DataFrame) -> pd.DataFrame:
    """Add relevance and priority fields to a validated legislative DataFrame."""
    processed = df.copy()
    processed["relevance_score"] = processed.apply(compute_relevance_score, axis=1)
    processed["priority"] = processed["relevance_score"].apply(assign_priority)
    processed["days_since_last_action"] = processed["last_action_date"].apply(_days_since_last_action)
    return processed


def _priority_explanation(row: pd.Series) -> str:
    reasons: list[str] = []
    keywords = _keyword_set(row.get("keywords"))
    keywords_text = " ".join(keywords)
    title_summary = f"{_safe_text(row.get('bill_title'))} {_safe_text(row.get('summary_text'))}"
    full_policy_text = f"{title_summary} {keywords_text}"
    committee = _safe_text(row.get("committee"))
    status = _safe_text(row.get("status"))

    if keywords.intersection(PORT_KEYWORDS):
        reasons.append("port, maritime, freight, or infrastructure keyword match")
    if _contains_any(title_summary, TITLE_SUMMARY_TERMS):
        reasons.append("direct transportation, funding, authorization, or workforce relevance")
    if _contains_any(committee, COMMITTEE_TERMS):
        reasons.append("relevant federal committee jurisdiction")
    if _contains_any(full_policy_text, ["wrda", "water infrastructure"]):
        reasons.append("WRDA or water infrastructure signal")
    if _contains_any(full_policy_text, ["new york", "new jersey", "panynj", "port authority"]):
        reasons.append("direct New York, New Jersey, or PANYNJ impact")
    if status.lower() in ACTIVE_STATUS_TERMS:
        reasons.append("active legislative movement")

    return "; ".join(reasons) if reasons else "general monitoring relevance"


def generate_briefing(df: pd.DataFrame) -> str:
    """Generate and write a daily policy briefing text report."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    processed = df.copy()
    high_priority = processed[processed["relevance_score"] >= 80].sort_values(
        by="last_action_date", ascending=False
    )
    watchlist = processed[
        (processed["status"].str.lower() == "stalled")
        & (processed["relevance_score"] >= 50)
    ].sort_values(by="relevance_score", ascending=False)

    lines = [
        "Daily Legislative Briefing",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "Executive Summary",
        f"- Total bills tracked: {len(processed)}",
        f"- High priority bills: {len(high_priority)}",
        f"- Average relevance score: {processed['relevance_score'].mean():.1f}",
        f"- Committees monitored: {processed['committee'].nunique()}",
        "",
        "High Priority Bills",
    ]

    if high_priority.empty:
        lines.append("- No high priority bills currently meet the alert threshold.")
    else:
        for _, row in high_priority.iterrows():
            date_label = row["last_action_date"].strftime("%Y-%m-%d") if pd.notna(row["last_action_date"]) else "Unknown date"
            lines.extend(
                [
                    f"- {row['bill_id']}: {row['bill_title']}",
                    f"  Score: {row['relevance_score']} | Status: {row['status']} | Last action: {date_label}",
                    f"  Explanation: {_priority_explanation(row)}.",
                    f"  URL: {row['url']}",
                ]
            )

    lines.extend(["", "Watchlist"])
    if watchlist.empty:
        lines.append("- No stalled but important bills are currently on the watchlist.")
    else:
        for _, row in watchlist.iterrows():
            lines.extend(
                [
                    f"- {row['bill_id']}: {row['bill_title']}",
                    f"  Score: {row['relevance_score']} | Committee: {row['committee']}",
                    f"  URL: {row['url']}",
                ]
            )

    briefing = "\n".join(lines) + "\n"
    BRIEFING_OUTPUT_PATH.write_text(briefing, encoding="utf-8")
    return briefing


def _summary_metrics(df: pd.DataFrame) -> pd.DataFrame:
    metrics = [
        {"metric": "Total Bills Tracked", "value": len(df)},
        {"metric": "High Priority Count", "value": int((df["priority"] == "HIGH PRIORITY").sum())},
        {"metric": "Medium Priority Count", "value": int((df["priority"] == "MEDIUM PRIORITY").sum())},
        {"metric": "Low Priority Count", "value": int((df["priority"] == "LOW PRIORITY").sum())},
        {"metric": "Average Relevance Score", "value": round(float(df["relevance_score"].mean()), 2)},
    ]
    summary = pd.DataFrame(metrics)
    committee_breakdown = (
        df.groupby("committee", dropna=False)
        .agg(bill_count=("bill_id", "count"), average_score=("relevance_score", "mean"))
        .reset_index()
    )
    committee_breakdown["average_score"] = committee_breakdown["average_score"].round(2)
    return pd.concat(
        [summary, pd.DataFrame([{"metric": "", "value": ""}]), committee_breakdown],
        ignore_index=True,
    )


def _excel_ready_df(df: pd.DataFrame) -> pd.DataFrame:
    """Order columns for analyst-friendly workbook viewing."""
    ordered_columns = [column for column in EXCEL_COLUMN_ORDER if column in df.columns]
    remaining_columns = [column for column in df.columns if column not in ordered_columns]
    return df[ordered_columns + remaining_columns].copy()


def _add_excel_table(worksheet, table_name: str) -> None:
    if worksheet.max_row < 2 or worksheet.max_column < 1:
        return
    ref = f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _style_excel_sheet(worksheet, *, table_name: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(bottom=Side(style="thin", color="D9E2F3"))
    hyperlink_font = Font(color="0563C1", underline="single")

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    header_lookup = {
        worksheet.cell(row=1, column=column_index).value: column_index
        for column_index in range(1, worksheet.max_column + 1)
    }

    for column_index in range(1, worksheet.max_column + 1):
        header = worksheet.cell(row=1, column=column_index).value
        column_letter = get_column_letter(column_index)
        worksheet.column_dimensions[column_letter].width = EXCEL_COLUMN_WIDTHS.get(str(header), 18)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
        worksheet.row_dimensions[row[0].row].height = 54

    for column_name in ("summary_text", "keywords"):
        column_index = header_lookup.get(column_name)
        if column_index:
            for cell in worksheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for item in cell:
                    item.alignment = Alignment(vertical="top", wrap_text=True)

    url_column = header_lookup.get("url")
    if url_column:
        for cell in worksheet.iter_cols(
            min_col=url_column,
            max_col=url_column,
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for item in cell:
                if item.value:
                    item.hyperlink = item.value
                    item.font = hyperlink_font

    date_column = header_lookup.get("last_action_date")
    if date_column:
        for cell in worksheet.iter_cols(
            min_col=date_column,
            max_col=date_column,
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for item in cell:
                item.number_format = "yyyy-mm-dd"

    score_column = header_lookup.get("relevance_score")
    if score_column:
        column_letter = get_column_letter(score_column)
        score_range = f"{column_letter}2:{column_letter}{worksheet.max_row}"
        worksheet.conditional_formatting.add(
            score_range,
            CellIsRule(operator="greaterThanOrEqual", formula=["80"], fill=PatternFill("solid", fgColor="F4CCCC")),
        )
        worksheet.conditional_formatting.add(
            score_range,
            CellIsRule(operator="between", formula=["50", "79"], fill=PatternFill("solid", fgColor="FCE5CD")),
        )
        worksheet.conditional_formatting.add(
            score_range,
            CellIsRule(operator="lessThan", formula=["50"], fill=PatternFill("solid", fgColor="D9EAD3")),
        )

    priority_column = header_lookup.get("priority")
    if priority_column:
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=priority_column)
            value = str(cell.value or "").upper()
            if value == "HIGH PRIORITY":
                cell.fill = PatternFill("solid", fgColor="CC0000")
                cell.font = Font(color="FFFFFF", bold=True)
            elif value == "MEDIUM PRIORITY":
                cell.fill = PatternFill("solid", fgColor="F6B26B")
                cell.font = Font(color="000000", bold=True)
            elif value == "LOW PRIORITY":
                cell.fill = PatternFill("solid", fgColor="B6D7A8")
                cell.font = Font(color="000000", bold=True)

    if table_name:
        _add_excel_table(worksheet, table_name)


def _style_summary_sheet(worksheet) -> None:
    _style_excel_sheet(worksheet)
    worksheet.freeze_panes = "A2"
    for row_index in range(2, worksheet.max_row + 1):
        if not worksheet.cell(row=row_index, column=1).value and not worksheet.cell(row=row_index, column=2).value:
            worksheet.row_dimensions[row_index].height = 12


def export_excel(df: pd.DataFrame) -> Path:
    """Export raw data, high priority bills, and summary metrics to Excel."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook_df = _excel_ready_df(df)
    high_priority = _excel_ready_df(df[df["priority"] == "HIGH PRIORITY"].copy())
    summary_metrics = _summary_metrics(df)

    with pd.ExcelWriter(EXCEL_OUTPUT_PATH, engine="openpyxl") as writer:
        workbook_df.to_excel(writer, sheet_name="Raw Data", index=False)
        high_priority.to_excel(writer, sheet_name="High Priority Bills", index=False)
        summary_metrics.to_excel(writer, sheet_name="Summary Metrics", index=False)

        workbook = writer.book
        _style_excel_sheet(workbook["Raw Data"], table_name="RawDataTable")
        _style_excel_sheet(workbook["High Priority Bills"], table_name="HighPriorityBillsTable")
        _style_summary_sheet(workbook["Summary Metrics"])
        workbook.active = workbook["High Priority Bills"] if len(high_priority) else workbook["Raw Data"]

    return EXCEL_OUTPUT_PATH


def run_tracker() -> pd.DataFrame:
    """Load, process, and export all tracker artifacts."""
    df = load_data()
    processed = process_legislation(df)
    export_excel(processed)
    generate_briefing(processed)
    return processed


if __name__ == "__main__":
    try:
        result = run_tracker()
        print(f"Processed {len(result)} bills.")
        print(f"Excel output: {EXCEL_OUTPUT_PATH}")
        print(f"Briefing output: {BRIEFING_OUTPUT_PATH}")
    except Exception as error:
        log_error(str(error))
        raise
